"""
Test suite for Cisco Asset Warranty Hardware UUID Validation and Deduplication.

Validates:
- Hardware Serial Number as immutable device UUID.
- Detection of 2 or more identical devices across multiple ingested files (CSV, HTML, JSON, TXT).
- Intelligent attribute merging/enrichment across duplicate occurrences.
- Discrepancy & conflict logging (differing hostnames, differing IPs).
- Statistical integrity of model_counts and warranty_status_counts (no inflation).
- Multi-format audit exports (HTML audit table, XLSX 'Duplicates Audit' sheet, JSON duplicates_audit array).
"""

import json
import openpyxl
import pytest
from pathlib import Path
from utils.asset_warranty_consolidator import AssetWarrantyConsolidator, consolidate_reports


@pytest.fixture
def sample_duplicate_csv_files(tmp_path):
    """Creates two CSV files with overlapping serial numbers and complementary attributes."""
    input_dir = tmp_path / "input_reports"
    output_dir = tmp_path / "output_reports"
    input_dir.mkdir()
    output_dir.mkdir()

    # File 1: Has Hostname, IP, Serial, but lacks Model and EOL
    file1 = input_dir / "cisco_warranty_file1.csv"
    file1.write_text(
        "Host,Hostname,Product ID (PID),Serial Number,Software Version,Warranty Status,End of Life,End of Support\n"
        "10.10.1.1,core-switch-01,,FOC12345678,16.12.4,Covered,,\n"
        "10.10.1.2,access-switch-01,WS-C2960X-24TD-L,FOC87654321,15.2(7)E,Expired,2025-10-31,2026-10-31\n"
    )

    # File 2: Same Serial FOC12345678, has Model and EOL, slight hostname difference, different non-duplicate device
    file2 = input_dir / "cisco_warranty_file2.csv"
    file2.write_text(
        "Host,Hostname,Product ID (PID),Serial Number,Software Version,Warranty Status,End of Life,End of Support\n"
        "10.10.1.1,core-sw01,C9300-48P,FOC12345678,,Covered,2028-05-31,2029-05-31\n"
        "10.10.1.3,edge-router-01,ISR4331/K9,FOC99999999,17.3.3,Covered,2029-12-31,2030-12-31\n"
    )

    return {
        'input_dir': input_dir,
        'output_dir': output_dir,
        'file1': file1,
        'file2': file2,
    }


def test_uuid_dedup_detection_and_attribute_merging(sample_duplicate_csv_files):
    """Verify that identical Serial Numbers are deduplicated and attributes are enriched."""
    fixtures = sample_duplicate_csv_files
    consolidator = AssetWarrantyConsolidator(
        input_dir=str(fixtures['input_dir']),
        output_dir=str(fixtures['output_dir'])
    )

    result = consolidator.consolidate_reports()

    assert result['success'] is True
    # 4 raw rows ingested across 2 files
    assert result['raw_devices_count'] == 4
    # FOC12345678 is duplicate (2 occurrences), so 3 unique devices total
    assert result['total_devices'] == 3
    assert result['duplicate_serials_count'] == 1
    assert result['duplicate_occurrences_removed'] == 1

    # Check JSON export
    json_path = fixtures['output_dir'] / result['output_json']
    assert json_path.exists()
    with open(json_path, 'r', encoding='utf-8') as f:
        data = json.load(f)

    # Verify duplicates summary in JSON
    dup_sum = data['duplicates_summary']
    assert dup_sum['raw_devices_count'] == 4
    assert dup_sum['unique_devices_count'] == 3
    assert dup_sum['duplicate_serials_count'] == 1
    assert dup_sum['duplicate_occurrences_removed'] == 1

    # Verify duplicates audit list in JSON
    audit = data['duplicates_audit']
    assert len(audit) == 1
    dup_entry = audit[0]
    assert dup_entry['serial'] == 'FOC12345678'
    assert dup_entry['occurrences'] == 2
    assert len(dup_entry['source_files']) == 2
    assert 'cisco_warranty_file1.csv' in dup_entry['source_files']
    assert 'cisco_warranty_file2.csv' in dup_entry['source_files']
    # Hostname conflict tracked ('core-switch-01' vs 'core-sw01')
    assert any('Differing hostnames' in c for c in dup_entry['conflicts'])

    # Find the merged device in devices list
    merged = next((d for d in data['devices'] if d['serial_number'] == 'FOC12345678'), None)
    assert merged is not None
    # Attribute merging verification:
    # Model enriched from File 2
    assert merged['product_id'] == 'C9300-48P'
    # Software version preserved from File 1
    assert merged['software_version'] == '16.12.4'
    # EOL preserved/refreshed from local Cisco EOL DB (C9300-48P → 2030-12-01)
    assert merged['end_of_life'] == '2030-12-01'
    assert merged['end_of_support'] == '2033-12-01'
    assert merged.get('eol_enrichment_source') == 'cisco_eol_database'


def test_statistical_integrity_no_count_inflation(sample_duplicate_csv_files):
    """Ensure model_counts and warranty_status_counts are calculated on deduplicated devices."""
    fixtures = sample_duplicate_csv_files
    consolidator = AssetWarrantyConsolidator(
        input_dir=str(fixtures['input_dir']),
        output_dir=str(fixtures['output_dir'])
    )

    result = consolidator.consolidate_reports()
    assert result['success'] is True

    summary = result['summary']
    # Total devices must be 3
    assert summary['total_devices'] == 3

    # Total counts in model_counts must sum to 3, NOT 4
    total_model_instances = sum(summary['model_counts'].values())
    assert total_model_instances == 3

    # Total counts in warranty_status_counts must sum to 3, NOT 4
    total_warranty_instances = sum(summary['warranty_status_counts'].values())
    assert total_warranty_instances == 3


def test_multi_sheet_xlsx_audit_export(sample_duplicate_csv_files):
    """Validate that XLSX export contains both 'Master Inventory' and 'Duplicates Audit' sheets."""
    fixtures = sample_duplicate_csv_files
    consolidator = AssetWarrantyConsolidator(
        input_dir=str(fixtures['input_dir']),
        output_dir=str(fixtures['output_dir'])
    )

    result = consolidator.consolidate_reports()
    xlsx_path = fixtures['output_dir'] / result['output_xlsx']
    assert xlsx_path.exists()

    wb = openpyxl.load_workbook(str(xlsx_path))
    sheet_names = wb.sheetnames

    assert "Master Inventory" in sheet_names
    assert "Duplicates Audit" in sheet_names

    ws_master = wb["Master Inventory"]
    # Row 1 is header, Rows 2-4 are the 3 unique devices
    assert ws_master.max_row == 4

    ws_audit = wb["Duplicates Audit"]
    # Row 1 is header, Row 2 is the duplicate serial entry
    assert ws_audit.max_row == 2
    assert ws_audit.cell(row=2, column=1).value == "FOC12345678"
    assert ws_audit.cell(row=2, column=2).value == 2  # occurrences


def test_html_audit_table_rendering(sample_duplicate_csv_files):
    """Validate that HTML report contains the Duplicate Hardware Audit section and cards."""
    fixtures = sample_duplicate_csv_files
    consolidator = AssetWarrantyConsolidator(
        input_dir=str(fixtures['input_dir']),
        output_dir=str(fixtures['output_dir'])
    )

    result = consolidator.consolidate_reports()
    html_path = fixtures['output_dir'] / result['output_html']
    assert html_path.exists()

    content = html_path.read_text(encoding='utf-8')
    assert "Duplicate Hardware Audit" in content
    assert "FOC12345678" in content
    assert "Duplicates Merged" in content
    assert "Raw Rows" in content


def test_placeholder_serials_do_not_collide(tmp_path):
    """Devices with unknown or blank serials must not be falsely treated as duplicates of each other."""
    input_dir = tmp_path / "input"
    output_dir = tmp_path / "output"
    input_dir.mkdir()
    output_dir.mkdir()

    f1 = input_dir / "report1.csv"
    f1.write_text(
        "Host,Hostname,Product ID (PID),Serial Number,Software Version,Warranty Status,End of Life,End of Support\n"
        "10.1.1.1,router-a,ISR4451,UNKNOWN,17.3,Covered,,\n"
        "10.1.1.2,router-b,ISR4451,UNKNOWN,17.3,Covered,,\n"
        "10.1.1.3,router-c,ISR4451,N/A,17.3,Covered,,\n"
    )

    consolidator = AssetWarrantyConsolidator(input_dir=str(input_dir), output_dir=str(output_dir))
    result = consolidator.consolidate_reports()

    assert result['success'] is True
    # None of these 3 distinct routers should be collapsed into one another
    assert result['total_devices'] == 3
    assert result['duplicate_occurrences_removed'] == 0
