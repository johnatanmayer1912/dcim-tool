# ============================================================================
# tests/test_asset_warranty_quad_export.py
# ============================================================================
# PURPOSE:
#   Comprehensive unit and integration test suite for the Quad-Format Export
#   (HTML, CSV, JSON, XLSX) in utils/asset_warranty_consolidator.py.
# ============================================================================

import csv
import json
import os
import pytest
import openpyxl
from utils.asset_warranty_consolidator import consolidate_reports, AssetWarrantyConsolidator


@pytest.fixture
def sample_reports(tmp_path):
    """Create two sample warranty HTML reports to consolidate."""
    r1 = tmp_path / "report_site1.html"
    r1.write_text("""
    <html><body>
    <table>
        <tr><th>Host</th><th>Hostname</th><th>Product ID (PID)</th><th>Serial Number</th><th>Software Version</th><th>Warranty Status</th><th>End of Life</th><th>End of Support</th></tr>
        <tr><td>10.1.1.1</td><td>SW-CORE-01</td><td>WS-C3850-24T</td><td>FCW12345678</td><td>16.12.4</td><td>covered</td><td>2028-10-31</td><td>2029-10-31</td></tr>
        <tr><td>10.1.1.2</td><td>SW-ACCESS-01</td><td>C9300-48P</td><td>FOC87654321</td><td>17.03.04</td><td>expired</td><td>2024-01-01</td><td>2025-01-01</td></tr>
    </table>
    </body></html>
    """, encoding="utf-8")

    r2 = tmp_path / "report_site2.html"
    r2.write_text("""
    <html><body>
    <table>
        <tr><th>Host</th><th>Hostname</th><th>Product ID (PID)</th><th>Serial Number</th><th>Software Version</th><th>Warranty Status</th><th>End of Life</th><th>End of Support</th></tr>
        <tr><td>10.2.1.1</td><td>SW-DIST-01</td><td>N9K-C93180YC-EX</td><td>SAL99887766</td><td>9.3(8)</td><td>covered</td><td>2027-05-30</td><td>2028-05-30</td></tr>
        <tr><td>10.1.1.1</td><td>SW-CORE-01</td><td>WS-C3850-24T</td><td>FCW12345678</td><td>16.12.4</td><td>covered</td><td>2028-10-31</td><td>2029-10-31</td></tr>
    </table>
    </body></html>
    """, encoding="utf-8")

    return [str(r1), str(r2)]


def test_consolidate_reports_quad_format(sample_reports, tmp_path):
    """Ensure consolidate_reports generates all 4 formats: HTML, CSV, JSON, and XLSX."""
    out_dir = str(tmp_path / "output")
    res = consolidate_reports(file_paths=sample_reports, output_dir=out_dir)

    assert res["success"] is True
    assert res["total_devices"] == 3  # Deduplicated across the 2 files
    assert res["total_reports"] == 2

    # Verify all 4 file keys are present
    master_html = res.get("master_file")
    csv_file = res.get("csv_file")
    json_file = res.get("json_file")
    xlsx_file = res.get("xlsx_file")

    assert master_html and os.path.isfile(master_html)
    assert csv_file and os.path.isfile(csv_file)
    assert json_file and os.path.isfile(json_file)
    assert xlsx_file and os.path.isfile(xlsx_file)

    assert master_html.endswith(".html")
    assert csv_file.endswith(".csv")
    assert json_file.endswith(".json")
    assert xlsx_file.endswith(".xlsx")


def test_master_html_content(sample_reports, tmp_path):
    """Ensure Master HTML contains download links and table."""
    out_dir = str(tmp_path / "output_html")
    res = consolidate_reports(file_paths=sample_reports, output_dir=out_dir)
    html_content = open(res["master_file"], encoding="utf-8").read()

    assert "Master Consolidated Inventory" in html_content
    assert "Download Master CSV" in html_content
    assert "Download Master JSON" in html_content
    assert "Download Master XLSX" in html_content
    assert "SW-CORE-01" in html_content
    assert "SW-ACCESS-01" in html_content
    assert "SW-DIST-01" in html_content


def test_master_csv_content(sample_reports, tmp_path):
    """Ensure CSV has proper headers and 3 data rows."""
    out_dir = str(tmp_path / "output_csv")
    res = consolidate_reports(file_paths=sample_reports, output_dir=out_dir)
    with open(res["csv_file"], mode="r", encoding="utf-8") as f:
        reader = csv.reader(f)
        rows = list(reader)

    assert len(rows) == 4  # 1 header row + 3 unique device rows
    headers = rows[0]
    assert "Host" in headers
    assert "Hostname" in headers
    assert "Serial Number" in headers
    assert "Warranty Status" in headers


def test_master_json_content(sample_reports, tmp_path):
    """Ensure JSON has metadata and array of devices."""
    out_dir = str(tmp_path / "output_json")
    res = consolidate_reports(file_paths=sample_reports, output_dir=out_dir)
    with open(res["json_file"], mode="r", encoding="utf-8") as f:
        data = json.load(f)

    assert "summary" in data
    assert data["total_devices"] == 3
    assert data["summary"]["total_files"] == 2
    assert "devices" in data
    assert len(data["devices"]) == 3

    hostnames = {d.get("hostname") for d in data["devices"]}
    assert hostnames == {"SW-CORE-01", "SW-ACCESS-01", "SW-DIST-01"}


def test_master_xlsx_content_and_styling(sample_reports, tmp_path):
    """Ensure XLSX opens in openpyxl with expected worksheet, headers, and colors."""
    out_dir = str(tmp_path / "output_xlsx")
    res = consolidate_reports(file_paths=sample_reports, output_dir=out_dir)
    
    wb = openpyxl.load_workbook(res["xlsx_file"])
    ws = wb.active
    assert ws.title == "Master Inventory"

    # Row 1 is Header Block
    headers = [ws.cell(row=1, column=col).value for col in range(1, 9)]
    assert headers == [
        "Host",
        "Hostname",
        "Product ID (PID)",
        "Serial Number",
        "Software Version",
        "Warranty Status",
        "End of Life",
        "End of Support",
    ]

    # Check header styling
    header_cell = ws.cell(row=1, column=1)
    assert header_cell.font.bold is True
    # Fill color is dark slate (1E293B)
    assert header_cell.fill.start_color.rgb.upper().endswith("1E293B")

    # Check data rows (rows 2, 3, 4)
    row_count = ws.max_row
    assert row_count == 4

    # Warranty status coloring — consolidator re-enriches from Cisco EOL DB
    # (file "expired"/"covered" labels are replaced with lifecycle status)
    warranty_statuses = [ws.cell(row=r, column=6).value for r in range(2, 5)]
    assert set(warranty_statuses) == {"expired", "active", "warning"}

    # EOL column refreshed from DB for known PIDs
    eol_vals = [ws.cell(row=r, column=7).value for r in range(2, 5)]
    assert "2030-12-01" in eol_vals  # C9300-48P
    assert "2027-08-31" in eol_vals  # N9K-C93180YC-EX
    assert "2023-03-07" in eol_vals  # WS-C3850-24T
